"""
Contrôle indépendant des archives produites par EvMasque (tests/e2e/sortie/).

    uv run --no-project --with pypdf --with openpyxl --with pypdfium2 --with pillow python tests/e2e/verifier-sortie.py

Ne fait pas confiance à l'outil : rouvre chaque fichier avec d'autres bibliothèques et
cherche les valeurs d'origine du corpus fictif.
"""
import io
import json
import re
import sys
import unicodedata
import zipfile
from pathlib import Path

import shutil
import subprocess
import pypdfium2 as pdfium
from openpyxl import load_workbook
from pypdf import PdfReader

ICI = Path(__file__).resolve().parent
SORTIE = ICI / "sortie"
INTERDITS = [
    "dupont", "exemple hydraulique", "musterfirma", "jean martin", "claire durand", "weber",
    "fa-20417-b", "fa20417b", "gh-118-03", "kl-5502", "mn-7781", "dupont-meca", "03 87 12 34 56",
    "solidworks", "clients",
]


def norme(s: str) -> str:
    s = unicodedata.normalize("NFD", s.lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def fuites(texte: str) -> list[str]:
    t = norme(texte)
    compact = re.sub(r"[^a-z0-9]", "", t)
    out = []
    for m in INTERDITS:
        mc = re.sub(r"[^a-z0-9]", "", norme(m))
        if m in t or (len(mc) >= 5 and mc in compact):
            out.append(m)
    return out


def main() -> int:
    res = json.loads((SORTIE / "resultat.json").read_text(encoding="utf-8"))
    mdp = res["motDePasse"]
    erreurs: list[str] = []
    envoi = next(SORTIE.glob("*_A_ENVOYER*.zip"))
    garder = next(SORTIE.glob("*_A_GARDER.zip"))
    print("Archive à envoyer :", envoi.name)

    # Déchiffrement par un code indépendant de celui qui a chiffré (tests/e2e/dechiffrer.mjs)
    dechiffreur = str(ICI / "dechiffrer.mjs")
    mauvais = subprocess.run(["node", dechiffreur, str(envoi), "mauvais-mot-de-passe", str(SORTIE / "rejet")], capture_output=True, text=True)
    if mauvais.returncode == 0:
        erreurs.append("archive lisible sans le bon mot de passe")
    else:
        print("  mauvais mot de passe refusé : oui")
    dossier = SORTIE / "envoi"
    shutil.rmtree(dossier, ignore_errors=True)
    bon = subprocess.run(["node", dechiffreur, str(envoi), mdp, str(dossier)], capture_output=True, text=True)
    print("  " + bon.stdout.strip().splitlines()[0] if bon.returncode == 0 else "  ÉCHEC : " + bon.stderr[-400:])
    if bon.returncode != 0:
        erreurs.append("déchiffrement impossible avec le bon mot de passe")
        return 1

    for chemin in sorted(p for p in dossier.rglob("*") if p.is_file()):
        nom = chemin.relative_to(dossier).as_posix()
        octets = chemin.read_bytes()
        if True:
            if fuites(nom):
                erreurs.append(f"{nom} : nom de fichier {fuites(nom)}")
            if "correspondance" in nom:
                erreurs.append("la table de correspondance est dans le paquet à envoyer")
            if nom.endswith(".pdf"):
                r = PdfReader(io.BytesIO(octets))
                texte = "".join((p.extract_text() or "") for p in r.pages)
                meta = {k: str(v) for k, v in (r.metadata or {}).items()}
                print(f"  {nom} : {len(r.pages)} p., texte={len(texte.strip())} car., méta={meta}")
                if texte.strip():
                    erreurs.append(f"{nom} : contient du texte")
                if fuites(json.dumps(meta)):
                    erreurs.append(f"{nom} : métadonnées {fuites(json.dumps(meta))}")
                doc = pdfium.PdfDocument(octets)
                for i in range(len(doc)):
                    doc[i].render(scale=0.35).to_pil().save(SORTIE / f"rendu_{Path(nom).stem}_{i + 1}.png")
            elif nom.endswith((".step", ".stp", ".csv", ".html")):
                t = octets.decode("utf-8", "replace")
                f = fuites(t)
                print(f"  {nom} : {len(t)} car., fuites={f}")
                if f:
                    erreurs.append(f"{nom} : {f}")
            elif nom.endswith(".xlsx"):
                wb = load_workbook(io.BytesIO(octets))
                cellules = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)
                props = " ".join(str(getattr(wb.properties, k)) for k in ("creator", "lastModifiedBy", "title", "subject", "keywords", "description"))
                feuilles = " ".join(wb.sheetnames)
                zf = zipfile.ZipFile(io.BytesIO(octets))
                brut = " ".join(zf.read(n).decode("utf-8", "replace") for n in zf.namelist())
                f = fuites(cellules + " " + props + " " + feuilles) + fuites(brut)
                prix = [h for h in re.findall(r"(Taux horaire|Prix unitaire|Marge)", cellules)]
                print(f"  {nom} : feuilles={wb.sheetnames}, fuites={sorted(set(f))}, colonnes prix={prix}")
                if f:
                    erreurs.append(f"{nom} : {sorted(set(f))}")
                if prix and "collecte" not in nom:
                    erreurs.append(f"{nom} : colonnes de prix présentes {prix}")

    with zipfile.ZipFile(garder) as z:
        noms = z.namelist()
        print("Archive à garder :", noms)
        table = load_workbook(io.BytesIO(z.read("table-de-correspondance.xlsx")))
        valeurs = " ".join(str(c.value) for ws in table.worksheets for row in ws.iter_rows() for c in row if c.value)
        if "FA-20417-B" not in valeurs or "P-001" not in valeurs:
            erreurs.append("table de correspondance incomplète")

    print()
    print(f"Requêtes réseau pendant le parcours : {len(res['requetes'])}")
    if res["requetes"]:
        erreurs.append(f"requêtes réseau : {res['requetes']}")
    if erreurs:
        print("ÉCHEC :")
        for e in erreurs:
            print("  -", e)
        return 1
    print("CONTRÔLE INDÉPENDANT RÉUSSI : aucune valeur d'origine dans le paquet à envoyer.")
    return 0


if __name__ == "__main__":
    sys.exit(main())

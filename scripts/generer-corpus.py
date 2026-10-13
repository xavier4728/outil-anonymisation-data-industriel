"""
Génère un corpus de test ENTIÈREMENT FICTIF pour EvMasque (aucune donnée de prospect).

    uv run --no-project --with reportlab --with pillow --with openpyxl python scripts/generer-corpus.py

Produit corpus/Essai-fictif/ : 4 pièces, tous les formats de la V1, des noms de clients
dans les cartouches, hors des cartouches, dans les métadonnées, les en-têtes STEP, les
propriétés Excel, une feuille masquée et des colonnes de prix.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A3, landscape
from reportlab.pdfgen import canvas

RACINE = Path(__file__).resolve().parent.parent / "corpus" / "Essai-fictif"
POLICE = "C:/Windows/Fonts/arial.ttf"


def plan_vectoriel(chemin: Path):
    l, h = landscape(A3)
    c = canvas.Canvas(str(chemin), pagesize=(l, h))
    c.setAuthor("Jean Martin")
    c.setTitle("FA-20417-B Bride - MECANIQUE DUPONT SAS")
    c.setSubject("Plan client DUPONT")
    c.setKeywords("DUPONT, FA-20417-B")
    m = 28
    c.rect(m, m, l - 2 * m, h - 2 * m)
    # la pièce
    c.circle(420, 470, 150)
    c.circle(420, 470, 60)
    c.line(220, 470, 620, 470)
    c.setFont("Helvetica", 14)
    c.drawString(450, 640, "Ø300 h7")
    c.drawString(440, 470 + 8, "Ø120 H8")
    c.drawString(640, 380, "Ra 0,8")
    c.drawString(80, 200, "Tolérances générales ISO 2768-mK")
    c.drawString(80, 175, "Selon norme interne DUPONT DN-114 rév. 3")
    c.saveState()
    c.translate(200, 470)
    c.rotate(90)
    c.drawString(-40, 0, "45 ±0,02")
    c.restoreState()
    # cartouche (bas droite, 190 x 60 mm)
    cw, ch = 190 * 2.835, 60 * 2.835
    x0, y0 = l - m - cw, m
    c.rect(x0, y0, cw, ch)
    for k in range(1, 4):
        c.line(x0, y0 + k * ch / 4, x0 + cw, y0 + k * ch / 4)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(x0 + 8, y0 + 3 * ch / 4 + 14, "Client : MÉCANIQUE DUPONT SAS")
    c.setFont("Helvetica", 11)
    c.drawString(x0 + 8, y0 + 2 * ch / 4 + 14, "N° de plan : FA-20417-B   Indice C   Échelle 1:1")
    c.drawString(x0 + 8, y0 + ch / 4 + 14, "Dessiné par : Jean Martin   j.martin@dupont-meca.fr")
    c.drawString(x0 + 8, y0 + 14, "Tél. 03 87 12 34 56 · Propriété de DUPONT SAS, reproduction interdite")
    c.showPage()
    c.save()


def image_plan(largeur, hauteur, lignes, cartouche, dpi=300):
    im = Image.new("L", (largeur, hauteur), 255)
    d = ImageDraw.Draw(im)
    f = ImageFont.truetype(POLICE, int(dpi * 0.16))
    fg = ImageFont.truetype(POLICE, int(dpi * 0.2))
    d.rectangle([60, 60, largeur - 60, hauteur - 60], outline=0, width=4)
    d.ellipse([largeur // 3, hauteur // 4, largeur // 3 + 700, hauteur // 4 + 700], outline=0, width=5)
    for i, t in enumerate(lignes):
        d.text((200, hauteur // 2 + i * 90), t, fill=0, font=f)
    cw, chh = int(190 / 25.4 * dpi), int(60 / 25.4 * dpi)
    x0, y0 = largeur - 80 - cw, hauteur - 80 - chh
    d.rectangle([x0, y0, x0 + cw, y0 + chh], outline=0, width=4)
    for i, t in enumerate(cartouche):
        d.text((x0 + 30, y0 + 30 + i * (chh - 40) // len(cartouche)), t, fill=0, font=fg if i == 0 else f)
    return im


def plan_scanne_pdf(chemin: Path):
    im = image_plan(
        2480, 3508,
        ["Ø25 g6", "Chanfrein 1x45°", "Livré selon spécification Exemple Hydraulique EH-22"],
        ["Client : Société Exemple Hydraulique SARL", "N° de plan : GH-118-03  Indice A", "Vérifié par : Claire Durand"],
    )
    im.convert("1").save(chemin, "PDF", resolution=300)


def plan_tif(chemin: Path):
    p1 = image_plan(
        3508, 2480,
        ["Ø40 h6", "Härten 58+2 HRC", "Werksnorm Musterfirma MN 2201"],
        ["Kunde: Antriebe Musterfirma GmbH", "Zeichnungs-Nr.: KL-5502", "Gezeichnet: H. Weber"],
    ).convert("1")
    p2 = image_plan(
        3508, 2480,
        ["Detail A", "Ra 1,6"],
        ["Kunde: Antriebe Musterfirma GmbH", "Zeichnungs-Nr.: KL-5502 Blatt 2", "info@musterfirma-antriebe.de"],
    ).convert("1")
    p1.save(chemin, compression="group4", dpi=(300, 300), save_all=True, append_images=[p2],
            tiffinfo={270: "Plan Antriebe Musterfirma GmbH", 315: "H. Weber"})


def photo_jpg(chemin: Path):
    im = image_plan(
        1800, 1273,
        ["Ø60", "Taraudage M8"],
        ["Client : MÉCANIQUE DUPONT SAS", "Réf. MN-7781"],
        dpi=150,
    ).convert("RGB")
    im.save(chemin, quality=88)


STEP = """ISO-10303-21;
HEADER;
/* SolidWorks export for {client} */
FILE_DESCRIPTION(('{ref} {client}'),'2;1');
FILE_NAME('C:\\\\Clients\\\\{dossier}\\\\{ref}.STEP','2026-03-12T10:22:41',('{auteur}'),('{client}'),
  'SwSTEP 2.0','SolidWorks 2022','');
FILE_SCHEMA(('AUTOMOTIVE_DESIGN {{ 1 0 10303 214 1 1 1 1 }}'));
ENDSEC;
DATA;
#1=APPLICATION_PROTOCOL_DEFINITION('international standard','automotive_design',2000,#2);
#2=APPLICATION_CONTEXT('core data for automotive mechanical design processes');
#10=PRODUCT('{ref}','{ref} pour {client}','',(#11));
#11=PRODUCT_CONTEXT('',#2,'mechanical');
#12=PRODUCT_DEFINITION_FORMATION('C','',#10);
#13=PRODUCT_DEFINITION('design','',#12,#14);
#20=PERSON('A1','{auteur}','',$,$,$);
#21=ORGANIZATION('O1','{client}','');
#30=CARTESIAN_POINT('',(0.,25.4,-12.5));
#31=CARTESIAN_POINT('',(1.E-3,2.5E2,0.));
#32=DIRECTION('',(0.,0.,1.));
#40=(LENGTH_UNIT()NAMED_UNIT(*)SI_UNIT(.MILLI.,.METRE.));
#50=DESCRIPTIVE_REPRESENTATION_ITEM('materiau','42CrMo4');
#51=SHAPE_REPRESENTATION('{ref}',(#30,#31),#2);
ENDSEC;
END-ISO-10303-21;
"""


def etude_xlsx(chemin: Path, client: str, ref: str, auteur: str):
    wb = Workbook()
    wb.properties.creator = auteur
    wb.properties.lastModifiedBy = auteur
    wb.properties.title = f"Etude {ref} {client}"
    ws = wb.active
    ws.title = "Etude"
    ws.append(["Client :", client])
    ws.append(["Contact", "j.martin@dupont-meca.fr"])
    ws.append([])
    ws.append(["Opération", "Machine", "Temps réglage (min)", "Temps cycle (s)", "Taux horaire (€/h)", "Prix unitaire", "Marge", "Temps total (s)"])
    ws.append(["Tournage ébauche", "Tornos DT26", 45, 38.5, 68, 4.2, 0.25, 120])
    ws.append([f"Fraisage {ref}", "Robodrill", 30, 22, 72, 3.1, 0.25, 80])
    ws["A5"].comment = Comment(f"Demande de {client}", auteur)
    cachee = wb.create_sheet("Tarifs clients")
    cachee.append([client, "remise 12 %"])
    cachee.sheet_state = "hidden"
    wb.save(chemin)


def main():
    pieces = {
        "FA-20417-B": ("MÉCANIQUE DUPONT SAS", "Jean Martin", "Dupont"),
        "GH-118-03": ("Société Exemple Hydraulique SARL", "Claire Durand", "ExempleHydraulique"),
        "KL-5502": ("Antriebe Musterfirma GmbH", "H. Weber", "Musterfirma"),
        "MN-7781": ("MÉCANIQUE DUPONT SAS", "Jean Martin", "Dupont"),
    }
    for ref in pieces:
        (RACINE / ref).mkdir(parents=True, exist_ok=True)
    plan_vectoriel(RACINE / "FA-20417-B" / "FA-20417-B_plan.pdf")
    plan_scanne_pdf(RACINE / "GH-118-03" / "GH-118-03 scan.pdf")
    plan_tif(RACINE / "KL-5502" / "KL-5502.tif")
    photo_jpg(RACINE / "MN-7781" / "photo plan MN-7781.jpg")
    for ref, (client, auteur, dossier) in pieces.items():
        (RACINE / ref / f"{ref}.step").write_text(STEP.format(ref=ref, client=client, auteur=auteur, dossier=dossier), encoding="utf-8")
    etude_xlsx(RACINE / "FA-20417-B" / "etude FA-20417-B.xlsx", *pieces["FA-20417-B"][:1], "FA-20417-B", "Jean Martin")
    (RACINE / "KL-5502" / "temps reels KL-5502.csv").write_text(
        "OF;Kunde;Zeichnungs-Nr.;Operation;Temps réel (s)\n"
        "OF2611;Antriebe Musterfirma GmbH;KL-5502;Tournage;41,2\n"
        "OF2611;Antriebe Musterfirma GmbH;KL-5502;Fraisage;19,8\n",
        encoding="utf-8",
    )
    (RACINE / "notes.docx").write_bytes(b"fichier non pris en charge")
    print("Corpus fictif ->", RACINE)


if __name__ == "__main__":
    main()
